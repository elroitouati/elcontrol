# -*- coding: utf-8 -*-
"""שקף השוואה: 5 צמדי גופנים לבחירה"""
import os
from pptx import Presentation
from pptx.util import Inches
import build as B

OPTIONS = [
    ("1", "עדין ומודרני", "IBM Plex Sans Hebrew ExtraLight", "IBM Plex Sans Hebrew Light", "IBM Plex Hebrew"),
    ("2", "ספרותי", "Frank Ruhl Libre Light", "Assistant Light", "Frank Ruhl Libre + Assistant"),
    ("3", "אלגנטי", "Bellefair", "Assistant Light", "Bellefair + Assistant"),
    ("4", "רך ועגול", "Rubik Light", "Rubik Light", "Rubik"),
    ("5", "כתב יד", "Solitreo", "Assistant Light", "Solitreo + Assistant"),
]

D = B.Deck("night")
R, P = D.R, D.P
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(B.W), Inches(B.H)
s = prs.slides.add_slide(prs.slide_layouts[6])
B.set_bg(s, B.make_bg("night", "fontsheet", 0.55))
D.text(s, B.M, 0.45, B.W - 2 * B.M, 0.6, [P(R("איזה גופן?", 30, D.TEXT, False, "Assistant Light"))])
cw, gap = 2.25, 0.22
for i, (num, label, tf, bf, names) in enumerate(OPTIONS):
    x = B.W - B.M - cw - i * (cw + gap)
    D.text(s, x, 1.35, cw, 0.6, [P(R(num, 30, D.ACC, False, "Assistant ExtraLight"))])
    D.text(s, x, 1.95, cw, 0.4, [P(R(label, 15, D.MUTED, False, "Assistant SemiBold"))])
    D.text(s, x, 2.45, cw, 1.1, [P(R("סופרמן", 54, D.TEXT, False, tf), line=0.9)], anchor="m")
    D.text(s, x, 3.55, cw, 0.5, [P(R("עידן עמדי", 22, D.ACC, False, tf))])
    D.text(s, x, 4.35, cw, 1.6, [P(R("שרק מי שצעד באפלה ובחושך", 17, D.TEXT, False, bf), line=1.1),
                                 P(R("יודע באמת מהו אור", 17, D.TEXT, False, bf), line=1.1)])
    D.text(s, x, 6.6, cw, 0.4, [P(R(names, 10, D.MUTED, False, "Assistant Light"), ltr=False)])
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "font_options.pptx")
prs.save(out); print(out)
