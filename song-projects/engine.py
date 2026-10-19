# -*- coding: utf-8 -*-
"""מנוע בניית מצגות הערכה חלופית בספרות (RTL מלא).
כל עבודה = תוכן (songs.py) + סגנון. הסגנון קובע רקע, גופנים, צבעים
ואת וריאנט הפריסה של כל שקף, כך שאין שתי עבודות שנראות אותו דבר.
"""
import math, os, random
from PIL import Image, ImageDraw, ImageFilter, ImageEnhance, ImageChops, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
W, H = 13.333, 7.5
M = 0.7
PW, PH = 1920, 1080


def rgb(h):
    return RGBColor.from_string(h)


def hex2t(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


# ======================= רקעים =======================
def _grain(img, amt=0.04, seed=3):
    random.seed(seed)
    n = Image.effect_noise(img.size, 9).convert("RGB")
    n = ImageEnhance.Brightness(n).enhance(amt)
    if sum(img.resize((1, 1)).getpixel((0, 0))) > 380:
        return ImageChops.subtract(img, n)
    return ImageChops.add(img, n)


def _radial(size, cx, cy, r, color, power=2.2, amt=1.0):
    g = Image.new("RGB", size, (0, 0, 0))
    d = ImageDraw.Draw(g)
    for rr in range(int(r), 0, -10):
        t = 1 - rr / r
        k = (t ** power) * amt
        d.ellipse([cx - rr, cy - rr, cx + rr, cy + rr], fill=tuple(int(c * k) for c in color))
    return g


def _photo_cover(path, w, h, focus=0.4):
    ph = Image.open(path).convert("RGB")
    r = max(w / ph.width, h / ph.height)
    ph = ph.resize((int(ph.width * r) + 1, int(ph.height * r) + 1), Image.LANCZOS)
    x = (ph.width - w) // 2
    y = int((ph.height - h) * focus)
    return ph.crop((x, y, x + w, y + h))


def make_bg(st, name, level=0.5, photo=None, photo_mode="blur", out_dir=None):
    """level 0..1: התקדמות בתוך המצגת (משמש לסגנונות שמשתנים לאורך העבודה)."""
    kind = st["bg_kind"]
    base = hex2t(st["bg"])
    img = Image.new("RGB", (PW, PH), base)
    acc = hex2t(st.get("glow", st["accent"]))

    if photo:
        ph = _photo_cover(photo, PW, PH, st.get("photo_focus", 0.4))
        if st.get("gray_photos"):
            ph = ImageOps.grayscale(ph).convert("RGB")
        if photo_mode == "blur":
            ph = ph.filter(ImageFilter.GaussianBlur(16))
        light_bg = sum(base) > 450
        if light_bg:
            img = Image.blend(ph, img, 0.78)
        else:
            img = Image.blend(ImageEnhance.Brightness(ph).enhance(0.55), img, 0.35)
        if st.get("photo_tint"):
            img = ImageOps.colorize(ImageOps.grayscale(img), black=hex2t(st["photo_tint"][0]), white=hex2t(st["photo_tint"][1]))

    if kind == "glow":
        cx, cy = st.get("glow_pos", (PW + 60, -140))
        img = ImageChops.add(img, _radial((PW, PH), cx, cy, 1650, acc, 2.3, 0.35 + 0.55 * level))
    elif kind == "horizon":  # זריחה שחורה: אופק בוער בתחתית השקף
        g = _radial((PW, PH), PW / 2, PH + 520, 1100, acc, 1.6, 0.35 + 0.5 * level)
        img = ImageChops.add(img, g)
        line = Image.new("L", (PW, PH), 0)
        ImageDraw.Draw(line).rectangle([0, PH - 6, PW, PH], fill=int(90 + 100 * level))
        line = line.filter(ImageFilter.GaussianBlur(6))
        img = ImageChops.add(img, Image.merge("RGB", [line.point(lambda v, c=c: int(v * c / 255)) for c in acc]))
    elif kind == "beam":  # מגדל של אור: קרן אחת שחוצה את החושך
        beam = Image.new("L", (PW, PH), 0)
        ox, oy = st.get("beam_origin", (PW * 0.08, PH * 0.92))
        a = math.radians(st.get("beam_angle", -18))
        L = 2600
        spread = 0.16
        p1 = (ox + L * math.cos(a - spread), oy + L * math.sin(a - spread))
        p2 = (ox + L * math.cos(a + spread), oy + L * math.sin(a + spread))
        ImageDraw.Draw(beam).polygon([(ox, oy), p1, p2], fill=int(26 + 30 * level))
        beam = beam.filter(ImageFilter.GaussianBlur(70))
        img = ImageChops.add(img, Image.merge("RGB", [beam.point(lambda v, c=c: int(v * c / 255)) for c in acc]))
        img = ImageChops.add(img, _radial((PW, PH), ox, oy, 420, acc, 2.6, 0.45))
    elif kind == "dusk":  # שחר במדבר: אינדיגו שמתבהר לחול
        top = hex2t(st["bg"])
        bot = hex2t(st["bg2"])
        grad = Image.new("RGB", (1, PH))
        for y in range(PH):
            t = (y / PH) ** (2.4 - 0.9 * level)
            grad.putpixel((0, y), tuple(int(top[i] + (bot[i] - top[i]) * t) for i in range(3)))
        img = grad.resize((PW, PH))
    elif kind == "field":  # אופק של שדה בתחתית
        band = Image.new("RGB", (PW, PH), base)
        d = ImageDraw.Draw(band)
        g1, g2 = hex2t(st["field1"]), hex2t(st["field2"])
        horizon = int(PH * (0.86 - 0.04 * level))
        for y in range(horizon, PH):
            t = (y - horizon) / (PH - horizon)
            d.line([(0, y), (PW, y)], fill=tuple(int(g1[i] + (g2[i] - g1[i]) * t) for i in range(3)))
        band = band.filter(ImageFilter.GaussianBlur(28))
        img = band
        img = ImageChops.add(img, _radial((PW, PH), PW * 0.82, horizon + 30, 900, (60, 45, 10), 2.4, 0.6 * level))
    elif kind == "lined":  # דף מחברת
        d = ImageDraw.Draw(img)
        line_c = hex2t(st["line"])
        for y in range(150, PH, 54):
            d.line([(0, y), (PW, y)], fill=line_c, width=2)
        mx = PW - 88
        d.line([(mx, 0), (mx, PH)], fill=hex2t(st["margin"]), width=3)
        d.line([(mx + 8, 0), (mx + 8, PH)], fill=hex2t(st["margin"]), width=1)
    elif kind == "concrete":
        n = Image.effect_noise((PW // 6, PH // 6), 40).convert("L").resize((PW, PH), Image.BICUBIC)
        n = n.filter(ImageFilter.GaussianBlur(3))
        tex = ImageOps.colorize(n, black=tuple(max(0, c - 18) for c in base), white=tuple(min(255, c + 10) for c in base))
        img = Image.blend(img, tex, 0.55)
    elif kind == "dots":  # יומן עם רשת נקודות עדינה
        d = ImageDraw.Draw(img)
        dc = hex2t(st["line"])
        for y in range(60, PH, 48):
            for x in range(60, PW, 48):
                d.ellipse([x - 2, y - 2, x + 2, y + 2], fill=dc)
    elif kind == "flag":  # לבן עם שני פסים דקים של הדגל
        if st.get("_stripes"):
            d = ImageDraw.Draw(img)
            c = hex2t(st["accent"])
            d.rectangle([0, 70, PW, 92], fill=c)
            d.rectangle([0, PH - 92, PW, PH - 70], fill=c)
    elif kind == "space":
        pass

    if kind in ("paper", "flag", "space", "lined", "dots", "field", "concrete"):
        vig = Image.new("L", (PW, PH), 0)
        ImageDraw.Draw(vig).ellipse([-PW * 0.25, -PH * 0.35, PW * 1.25, PH * 1.35], fill=255)
        vig = vig.filter(ImageFilter.GaussianBlur(240))
        img = Image.composite(img, ImageEnhance.Brightness(img).enhance(st.get("vignette", 0.95)), vig)
    img = _grain(img, st.get("grain", 0.04))
    out_dir = out_dir or os.path.join(HERE, "_bg", st["id"])
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, name + ".png")
    img.save(out, optimize=True)
    return out


def prep_photo(st, path, tag=""):
    """גרסת תמונה לפי הסגנון (שחור-לבן, טון צבע)."""
    if not (st.get("gray_photos") or st.get("tone_photos")):
        return path
    out_dir = os.path.join(HERE, "_bg", st["id"])
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "ph_" + tag + os.path.basename(path).rsplit(".", 1)[0] + ".jpg")
    im = Image.open(path).convert("RGB")
    if st.get("gray_photos"):
        im = ImageEnhance.Contrast(ImageOps.grayscale(im)).enhance(1.1).convert("RGB")
    if st.get("tone_photos"):
        a, b = st["tone_photos"]
        im = ImageOps.colorize(ImageOps.grayscale(im), black=hex2t(a), white=hex2t(b))
    im.save(out, quality=92)
    return out


# ======================= טקסט =======================
class Kit:
    def __init__(self, st):
        self.st = st
        self.TEXT, self.MUTED, self.ACC = rgb(st["text"]), rgb(st["muted"]), rgb(st["accent"])
        self.ACC2 = rgb(st.get("accent2", st["accent"]))
        self.HAIR = rgb(st["hair"])
        self.PANEL = rgb(st.get("panel", st["bg"]))
        self.F = st["fonts"]

    def R(self, text, size=22, color=None, font="body", bold=False):
        if font in ("display", "title") and self.st.get("display_bold"):
            bold = True
        f = self.F.get(font, font)
        return (text, dict(size=size, color=color or self.TEXT, font=f, bold=bold))

    @staticmethod
    def P(*runs, align="r", after=None, line=None, ltr=False):
        return dict(runs=list(runs), align=align, after=after, line=line, ltr=ltr)

    def text(self, slide, x, y, w, h, paras, anchor="t", name=None, rot=0, wrap=True):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        if rot:
            tb.rotation = rot
        tf = tb.text_frame
        tf.word_wrap = wrap
        bp = tf._txBody.find(qn("a:bodyPr"))
        for k in ("lIns", "rIns", "tIns", "bIns"):
            bp.set(k, "0")
        bp.set("rtlCol", "1")
        bp.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pPr = para._p.get_or_add_pPr()
            pPr.set("algn", p["align"])
            pPr.set("rtl", "0" if p["ltr"] else "1")
            if p["line"]:
                ls = etree.SubElement(pPr, qn("a:lnSpc"))
                etree.SubElement(ls, qn("a:spcPct")).set("val", str(int(p["line"] * 100000)))
            if p["after"] is not None:
                sa = etree.SubElement(pPr, qn("a:spcAft"))
                etree.SubElement(sa, qn("a:spcPts")).set("val", str(int(p["after"] * 100)))
            etree.SubElement(pPr, qn("a:buNone"))
            for t, o in p["runs"]:
                r = para.add_run()
                r.text = t
                rPr = r._r.get_or_add_rPr()
                rPr.set("lang", "he-IL"); rPr.set("altLang", "en-US")
                rPr.set("sz", str(int(o["size"] * 100)))
                rPr.set("b", "1" if o["bold"] else "0")
                fill = etree.SubElement(rPr, qn("a:solidFill"))
                etree.SubElement(fill, qn("a:srgbClr")).set("val", str(o["color"]))
                for tag in ("a:latin", "a:cs"):
                    etree.SubElement(rPr, qn(tag)).set("typeface", o["font"])
            e = etree.SubElement(para._p, qn("a:endParaRPr"))
            e.set("lang", "he-IL")
            e.set("sz", str(int(p["runs"][0][1]["size"] * 100)) if p["runs"] else "2200")
        return tb

    def hair(self, slide, x, y, w, color=None, weight=0.75, vertical=False):
        x2, y2 = (x, y + w) if vertical else (x + w, y)
        s = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x2), Inches(y2))
        s.line.color.rgb = color or self.HAIR
        s.line.width = Pt(weight)
        stl = s._element.find(qn("p:style"))
        if stl is not None:
            s._element.remove(stl)
        return s

    def rect(self, slide, x, y, w, h, fill=None, line=None, line_w=1.0, alpha=None, dash=False, name=None):
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            s.name = name
        if fill is None:
            s.fill.background()
        else:
            s.fill.solid(); s.fill.fore_color.rgb = fill
            if alpha is not None:
                clr = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 1000)))
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(line_w)
            if dash:
                ln = s.line._get_or_add_ln()
                etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
        st = s._element.find(qn("p:style"))
        if st is not None:
            s._element.remove(st)
        s.text_frame.text = ""
        return s

    def dot(self, slide, x, y, d, fill):
        s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        s.fill.solid(); s.fill.fore_color.rgb = fill
        s.line.fill.background()
        st = s._element.find(qn("p:style"))
        if st is not None:
            s._element.remove(st)
        return s


def picture_fill(slide, path, x, y, w, h, focus=0.3, name=None):
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    br, ir = w / h, iw / ih
    if ir > br:
        c = 1 - br / ir
        fx = focus if isinstance(focus, float) and focus != 0.3 else 0.5
        pic.crop_left, pic.crop_right = c * 0.5, c * 0.5
    else:
        c = 1 - ir / br
        pic.crop_top, pic.crop_bottom = c * focus, c * (1 - focus)
    if name:
        pic.name = name
    return pic


def set_bg(slide, path):
    pic = slide.shapes.add_picture(path, 0, 0, Inches(W), Inches(H))
    pic.name = "Background"
    tree = slide.shapes._spTree
    tree.remove(pic._element)
    tree.insert(2, pic._element)


def add_transition(slide, kind="fade"):
    xml = {
        "fade": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:fade/></p:transition>',
        "push": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:push dir="l"/></p:transition>',
        "wipe": '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:wipe dir="l"/></p:transition>',
    }[kind]
    slide._element.find(qn("p:clrMapOvr")).addnext(etree.fromstring(xml))
