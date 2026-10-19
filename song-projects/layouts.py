# -*- coding: utf-8 -*-
"""ספריית פריסות. לכל סוג שקף כמה וריאנטים; הסגנון בוחר אחד."""
import os
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from engine import (Kit, W, H, M, make_bg, prep_photo, picture_fill, set_bg, add_transition, ASSETS, rgb)


def est_w(text, size, factor=0.5):
    """הערכת רוחב טקסט באינצ'ים."""
    return len(text) * factor * size / 72.0


class Builder:
    def __init__(self, song, st):
        self.s, self.st = song, st
        self.K = Kit(st)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0
        self.total = 10

    # ---------- תשתית ----------
    def asset(self, name):
        return os.path.join(ASSETS, name)

    def new(self, level, photo=None, stripes=False, photo_mode="blur"):
        self.n += 1
        sl = self.prs.slides.add_slide(self.blank)
        st = dict(self.st)
        st["_stripes"] = stripes
        set_bg(sl, make_bg(st, f"s{self.n:02d}", level, photo, photo_mode))
        add_transition(sl, self.st.get("transition", "fade"))
        return sl

    def R(self, *a, **k):
        return self.K.R(*a, **k)

    def P(self, *a, **k):
        return self.K.P(*a, **k)

    def title(self, sl, text, x=None, w=None, size=None, align="r", y=0.5):
        st = self.st
        x = M if x is None else x
        w = (W - 2 * M) if w is None else w
        self.K.text(sl, x, y, w, 0.95, [self.P(self.R(text, size or st.get("title_size", 42), self.K.TEXT, "title"), align=align)],
                    anchor="m", name="Title")

    def pageno(self, sl, x=None, color=None):
        fmt = self.st.get("pageno", "{n}")
        x = M if x is None else x
        self.K.text(sl, x, 6.95, 1.2, 0.3, [self.P(self.R(fmt.format(n=self.n, t=self.total, n2=f"{self.n:02d}"), 12, color or self.K.MUTED, "label"),
                                                   align="l", ltr=True)], anchor="m", name="Slide number")

    def hl_runs(self, line, hls, size, color, hl_color, font):
        """מחלק שורה לריצות ומדגיש מילים מתוך hls."""
        runs, i = [], 0
        while i < len(line):
            hit = None
            for h in hls:
                if h and line.startswith(h, i):
                    if hit is None or len(h) > len(hit):
                        hit = h
            if hit:
                runs.append(self.R(hit, size, hl_color, font))
                i += len(hit)
            else:
                j = i
                while j < len(line) and not any(h and line.startswith(h, j) for h in hls):
                    j += 1
                runs.append(self.R(line[i:j], size, color, font))
                i = j
        return runs

    def excerpt_paras(self, lines, hls, size, color=None, hl=None, font="quote", align="r", line=1.12):
        color = color or self.K.TEXT
        hl = hl or self.K.ACC
        return [self.P(*self.hl_runs(l, hls, size, color, hl, font), align=align, line=line, after=4) for l in lines]

    def credit(self, sl, x, y, w, text, color=None, align="l"):
        self.K.text(sl, x, y, w, 0.28, [self.P(self.R(text, 10, color or self.K.MUTED, "label"), align=align)], anchor="m")

    # =========================== שער ===========================
    def cover(self):
        s, st, K = self.s, self.st, self.K
        v = st["v"]["cover"]
        size = st.get("cover_size", 120)
        who = s.get("names")
        if isinstance(who, list) and v != "poster":
            who = " ו".join(who)
        names = self.P(self.R(who, 22, K.TEXT, "strong") if who else self.R("שמות המגישים", 20, K.MUTED, "body"))
        label = "הערכה חלופית בספרות · שכבה ט׳"
        if v == "center":
            sl = self.new(0.0, stripes=st.get("stripes_on_cover", False))
            K.text(sl, M, 0.85, W - 2 * M, 0.5, [self.P(self.R(label, 18, K.ACC, "label"), align="ctr")])
            K.text(sl, M, 1.55, W - 2 * M, 2.6, [self.P(self.R(s["title"], size, K.TEXT, "display"), align="ctr", line=0.9)], anchor="m", name="Title")
            K.text(sl, M, 4.25, W - 2 * M, 0.8, [self.P(self.R(s["artist"], 34, K.ACC, "light"), align="ctr")])
            K.text(sl, M, 5.05, W - 2 * M, 0.6, [self.P(self.R(s["tagline"], 22, K.MUTED, "body"), align="ctr")])
            names["align"] = "ctr"
            K.text(sl, M, 6.25, W - 2 * M, 0.5, [names])
        elif v == "right-stack":
            sl = self.new(0.0)
            K.text(sl, M, 0.75, W - 2 * M, 0.5, [self.P(self.R(label, 18, K.ACC, "label"))])
            K.text(sl, M, 1.5, W - 2 * M, 2.5, [self.P(self.R(s["title"], size, K.TEXT, "display"), line=0.9)], anchor="m", name="Title")
            K.text(sl, M, 4.1, W - 2 * M, 0.8, [self.P(self.R(s["artist"], 34, K.ACC, "light"))])
            K.text(sl, M, 4.9, W - 2 * M, 0.6, [self.P(self.R(s["tagline"], 22, K.MUTED, "body"))])
            K.text(sl, M, 6.2, W - 2 * M, 0.5, [names])
        elif v == "photo-split":
            sl = self.new(0.0)
            pw = 5.6
            picture_fill(sl, prep_photo(st, self.asset(s["cover_photo"])), 0, 0, pw, H, s.get("cover_focus", 0.35), name="Cover photo")
            self.credit(sl, 0.25, 7.12, pw - 0.4, s["cover_credit"], color=RGBColor(0xEE, 0xEE, 0xEE))
            x0 = pw + 0.7
            ww = W - M - x0
            K.text(sl, x0, 0.8, ww, 0.5, [self.P(self.R(label, 18, K.ACC, "label"))])
            K.text(sl, x0, 1.55, ww, 2.4, [self.P(self.R(s["title"], size, K.TEXT, "display"), line=0.9)], anchor="m", name="Title")
            K.text(sl, x0, 4.05, ww, 0.8, [self.P(self.R(s["artist"], 34, K.ACC, "light"))])
            K.text(sl, x0, 4.85, ww, 0.9, [self.P(self.R(s["tagline"], 22, K.MUTED, "body"), line=1.05)])
            K.text(sl, x0, 6.2, ww, 0.5, [names])
        elif v == "band":
            sl = self.new(0.0)
            K.text(sl, M, 0.55, W - 2 * M, 0.45, [self.P(self.R(label, 18, K.ACC, "label"))])
            K.text(sl, M, 0.95, W - 2 * M, 1.75, [self.P(self.R(s["title"], size, K.TEXT, "display"), line=0.9)], anchor="m", name="Title")
            picture_fill(sl, prep_photo(st, self.asset(s["cover_photo"])), 0, 2.9, W, 2.55, s.get("cover_focus", 0.5), name="Cover photo")
            self.credit(sl, M, 5.5, 6, s["cover_credit"])
            K.text(sl, 6.0, 5.7, W - M - 6.0, 0.7, [self.P(self.R(s["artist"], 32, K.ACC, "light"))])
            K.text(sl, M, 6.35, W - 2 * M, 0.5, [self.P(self.R(s["tagline"], 20, K.MUTED, "body"))])
            K.text(sl, M, 5.75, 5.0, 0.5, [self.P(self.R(s.get("names") or "שמות המגישים", 20, K.TEXT if s.get("names") else K.MUTED, "strong" if s.get("names") else "body"), align="l")])
        elif v == "poster":
            sl = self.new(0.0)
            K.text(sl, M, 0.6, W - 2 * M, 0.5, [self.P(self.R(label, 18, K.TEXT, "label"))], name="Label")
            who_lines = who if isinstance(who, list) else ([who] if who else None)
            if who_lines:
                K.text(sl, M, 0.55, 6.2, 0.4 * len(who_lines), [self.P(self.R(l, 16, K.TEXT, "strong"), align="l", line=1.05) for l in who_lines], name="Names")
            else:
                K.text(sl, M, 0.6, 5.0, 0.5, [self.P(self.R("שמות המגישים", 20, K.MUTED, "body"), align="l")])
            K.text(sl, M, 1.3, W - 2 * M, 0.8, [self.P(self.R(s["artist"], 34, K.ACC, "light"))])
            K.text(sl, M - 0.1, 2.0, W - 2 * M + 0.2, 4.6, [self.P(self.R(s["title"], size, K.TEXT, "display", bold=st.get("display_bold", False)), line=0.8)],
                   anchor="b", name="Title")
            K.text(sl, M, 6.7, W - 2 * M, 0.45, [self.P(self.R(s["tagline"], 20, K.MUTED, "body"))])

    # =========================== חלק א' ===========================
    def part_a(self):
        s, st, K = self.s, self.st, self.K
        v = st["v"]["part_a"]
        qsize = st.get("quote_size", 28)
        one = s.get("singular")
        picked = "הקטע שבחרתי" if one else "הקטע שבחרנו"
        why = [self.P(self.R("למה בחרתי בקטע" if one else "למה בחרנו בקטע", 22, K.ACC, "strong"), after=4),
               self.P(self.R(s["why"], 22), line=1.08)]
        pers = [self.P(self.R("החיבור האישי שלי" if one else "החיבור האישי שלנו", 22, K.ACC, "strong"), after=4),
                self.P(self.R(s["personal"], 22), line=1.08)]
        src = f"״{s['title']}״, {s['artist_short']}"
        if v == "quote-right":
            sl = self.new(0.1)
            self.title(sl, picked)
            K.text(sl, 7.25, 1.55, W - M - 7.25, 1.0, [self.P(self.R("״", 96, K.ACC, "display"), line=0.8)])
            K.text(sl, 7.25, 2.45, W - M - 7.25, 3.4, self.excerpt_paras(s["excerpt"], s["excerpt_hl"], qsize), anchor="t")
            K.text(sl, 7.25, 6.2, W - M - 7.25, 0.4, [self.P(self.R(src, 16, K.MUTED, "body"))])
            K.text(sl, M, 1.7, 6.0, 2.5, why)
            K.text(sl, M, 4.3, 6.0, 2.6, pers)
        elif v == "quote-top":
            sl = self.new(0.1)
            self.title(sl, picked)
            n = len(s["excerpt"])
            K.text(sl, M, 1.55, W - 2 * M, 0.55 * n, self.excerpt_paras(s["excerpt"], s["excerpt_hl"], qsize, line=1.0), anchor="t")
            K.text(sl, M, 1.55, 4.5, 0.4, [self.P(self.R(src, 16, K.MUTED, "body"), align="l")])
            yh = 1.55 + 0.52 * n + 0.25
            K.hair(sl, M, yh, W - 2 * M)
            K.text(sl, 6.95, yh + 0.3, W - M - 6.95, 6.95 - yh - 0.3, why)
            K.text(sl, M, yh + 0.3, 5.75, 6.95 - yh - 0.3, pers)
        elif v == "quote-left-panel":
            sl = self.new(0.1)
            pw = 6.4
            K.rect(sl, 0, 0, pw, H, fill=rgb(st["panel"]), name="Quote panel")
            ptxt = rgb(st.get("panel_text", st["text"]))
            phl = rgb(st.get("panel_hl", st["accent"]))
            K.text(sl, 0.6, 0.75, pw - 1.2, 0.5, [self.P(self.R(picked, 18, ptxt, "label"))])
            K.text(sl, 0.6, 1.4, pw - 1.2, 4.6, self.excerpt_paras(s["excerpt"], s["excerpt_hl"], st.get("panel_quote_size", qsize - 4), ptxt, phl), anchor="m")
            K.text(sl, 0.6, 6.4, pw - 1.2, 0.4, [self.P(self.R(src, 16, ptxt, "body"))])
            x0 = pw + 0.7
            self.title(sl, "למה דווקא הקטע הזה", x=x0, w=W - M - x0, size=st.get("title_size", 42) - 6)
            K.text(sl, x0, 1.75, W - M - x0, 2.5, why)
            K.text(sl, x0, 4.35, W - M - x0, 2.6, pers)
        self.pageno(sl)

    # =========================== האמן ===========================
    def artist(self):
        s, st, K = self.s, self.st, self.K
        v = st["v"]["artist"]
        sl = self.new(0.2)
        photo = prep_photo(st, self.asset(s["artist_photo"]))
        focus = s.get("artist_focus", 0.2)
        bio = s["bio"]

        def rows(x0, x1, y0, step=1.04, label_w=1.35):
            y = y0
            for i, (lab, t) in enumerate(bio):
                if i:
                    K.hair(sl, x0, y - 0.12, x1 - x0)
                K.text(sl, x1 - label_w, y, label_w, 0.5, [self.P(self.R(lab, 24, K.ACC, "light"))])
                K.text(sl, x0, y + 0.02, x1 - label_w - x0 - 0.25, 0.92, [self.P(self.R(t, 22), line=1.0)])
                y += step

        if v == "bleed-left":
            pw = 4.75
            picture_fill(sl, photo, 0, 0, pw, H, focus, name="Artist photo")
            self.credit(sl, 0.2, 7.12, pw - 0.3, s["artist_credit"], RGBColor(0xEE, 0xEE, 0xEE))
            self.title(sl, s["artist_name"], x=pw + 0.6, w=W - M - pw - 0.6)
            rows(pw + 0.6, W - M, 1.7)
            self.pageno(sl, x=pw + 0.6)
        elif v == "bleed-right":
            pw = 4.75
            picture_fill(sl, photo, W - pw, 0, pw, H, focus, name="Artist photo")
            self.credit(sl, W - pw + 0.2, 7.12, pw - 0.4, s["artist_credit"], RGBColor(0xEE, 0xEE, 0xEE), align="r")
            self.title(sl, s["artist_name"], x=M, w=W - pw - 0.6 - M)
            rows(M, W - pw - 0.6, 1.7)
            self.pageno(sl)
        elif v == "framed-left":
            pw, ph_ = 4.0, 5.0
            picture_fill(sl, photo, M, 1.55, pw, ph_, focus, name="Artist photo")
            self.credit(sl, M, 1.55 + ph_ + 0.08, pw, s["artist_credit"])
            self.title(sl, s["artist_name"])
            rows(M + pw + 0.6, W - M, 1.7)
            self.pageno(sl)
        elif v == "band-top":
            bh = 2.75
            picture_fill(sl, photo, 0, 0, W, bh, focus, name="Artist photo")
            self.credit(sl, M, bh + 0.05, 7, s["artist_credit"])
            self.title(sl, s["artist_name"], y=bh + 0.2, size=st.get("title_size", 42) - 4)
            colw = (W - 2 * M - 0.6) / 2
            for i, (lab, t) in enumerate(bio):
                col = 0 if i < 3 else 1
                row = i if i < 3 else i - 3
                x = W - M - colw - col * (colw + 0.6)
                y = 4.15 + row * 1.0
                K.text(sl, x, y, colw, 0.95, [self.P(self.R(lab + "  ", 22, K.ACC, "strong"), self.R(t, 22), line=1.0)])
            self.pageno(sl)

    # =========================== מילות השיר ===========================
    def lyrics(self):
        s, st, K = self.s, self.st, self.K
        stanzas = s.get("lyrics")
        if not stanzas:
            sl = self.new(0.3)
            self.title(sl, "מילות השיר")
            K.rect(sl, M, 1.6, W - 2 * M, 4.85, line=K.HAIR, line_w=1.25, dash=True, name="Lyrics placeholder")
            K.text(sl, M, 1.6, W - 2 * M, 4.85, [
                self.P(self.R("כאן יודבקו מילות השיר המלאות", 26, K.MUTED, "light"), align="ctr", after=8),
                self.P(self.R("לפי חלוקת הבתים והשורות המקורית", 18, K.MUTED, "body"), align="ctr")], anchor="m", name="Lyrics")
            K.text(sl, M, 6.6, W - 2 * M, 0.35, [self.P(self.R(s["credits"], 14, K.MUTED, "body"))], anchor="m")
            self.pageno(sl)
            return
        size = st.get("lyrics_size", 20)
        wrap_at = st.get("lyrics_wrap", 36)
        cap = 15.5  # שורות לעמודה
        from collections import Counter
        seen = Counter()
        for stz in stanzas:
            for ln in set(stz):
                seen[ln] += 1
        chorus = {ln for ln, c in seen.items() if c >= 2}
        cost = lambda ln: 2 if len(ln) > wrap_at else 1
        cols, cur, used = [], [], 0.0
        for stz in stanzas:
            need = sum(cost(l) for l in stz)
            if cur and used + 0.6 + need > cap:
                cols.append(cur); cur, used = [], 0.0
            if need > cap:  # בית ארוך מעמודה: מפצלים
                part, pu = [], 0
                for l in stz:
                    if part and pu + cost(l) > cap:
                        cur.append(part); cols.append(cur); cur, used, part, pu = [], 0.0, [], 0
                    part.append(l); pu += cost(l)
                cur.append(part); used = pu
                continue
            cur.append(stz); used += (0.6 if len(cur) > 1 else 0) + need
        if cur:
            cols.append(cur)
        pages = [cols[i:i + 2] for i in range(0, len(cols), 2)]
        for pi, pcols in enumerate(pages):
            sl = self.new(0.3 + 0.03 * pi)
            self.title(sl, "מילות השיר" if pi == 0 else "מילות השיר (המשך)")
            colw = (W - 2 * M - 0.7) / 2
            for ci, col in enumerate(pcols):
                x = W - M - colw - ci * (colw + 0.7)
                paras = []
                for si, stz in enumerate(col):
                    for li, ln in enumerate(stz):
                        ch = ln in chorus
                        paras.append(self.P(self.R(ln, size, K.ACC if ch else K.TEXT, "quote"), line=1.0,
                                            after=10 if (li == len(stz) - 1 and si < len(col) - 1) else 0))
                K.text(sl, x, 1.6, colw, 5.1, paras, name=f"Lyrics column {ci + 1}")
            if pi == len(pages) - 1:
                K.text(sl, M, 6.72, W - 2 * M, 0.3, [self.P(self.R(s["credits"], 12, K.MUTED, "body"))], anchor="m")
            self.pageno(sl)

    # =========================== מושגים ===========================
    def concept(self, i):
        s, st, K = self.s, self.st, self.K
        c = s["concepts"][i]
        v = st["v"]["concept"]
        sl = self.new(0.4 + 0.05 * i)
        lab = ["מושג ראשון", "מושג שני"][i]
        qs = st.get("cquote_size", 30)
        how = [self.P(self.R("איך זה מופיע בשיר", 22, K.ACC, "strong"), after=4), self.P(self.R(c["how"], 22), line=1.04)]
        why = [self.P(self.R("מה זה מוסיף למשמעות", 22, K.ACC, "strong"), after=4), self.P(self.R(c["why"], 22), line=1.04)]
        if v == "split":
            K.text(sl, 8.6, 0.6, W - M - 8.6, 0.4, [self.P(self.R(lab, 18, K.MUTED, "strong"))])
            K.text(sl, 8.6, 1.0, W - M - 8.6, 1.2, [self.P(self.R(c["name"], 56, K.ACC, "display"))], anchor="m", name="Title")
            K.text(sl, 8.6, 2.35, W - M - 8.6, 1.95, [self.P(self.R("הגדרה: ", 22, K.MUTED, "strong"), self.R(c["definition"], 22), line=1.05)])
            K.text(sl, M, 0.75, 7.4, 1.0, [self.P(self.R("״", 90, K.ACC, "display"), line=0.8)])
            K.text(sl, M, 1.75, 7.4, 2.5, self.excerpt_paras(c["quote"], c["hl"], qs), anchor="t")
            K.hair(sl, M, 4.5, W - 2 * M)
            K.text(sl, 6.95, 4.75, W - M - 6.95, 2.15, how)
            K.text(sl, M, 4.75, 5.75, 2.15, why)
        elif v == "stacked":
            K.text(sl, M, 0.5, W - 2 * M, 0.4, [self.P(self.R(lab, 18, K.MUTED, "strong"))])
            K.text(sl, 6.6, 0.85, W - M - 6.6, 1.1, [self.P(self.R(c["name"], 54, K.ACC, "display"))], anchor="m", name="Title")
            K.text(sl, M, 0.95, 5.6, 1.1, [self.P(self.R(c["definition"], 20, K.MUTED, "body"), align="l", line=1.04)], anchor="m")
            K.hair(sl, M, 2.2, W - 2 * M)
            K.text(sl, M, 2.4, W - 2 * M, 1.85, self.excerpt_paras(c["quote"], c["hl"], qs, align="ctr"), anchor="m")
            K.hair(sl, M, 4.45, W - 2 * M)
            K.text(sl, 6.95, 4.7, W - M - 6.95, 2.2, how)
            K.text(sl, M, 4.7, 5.75, 2.2, why)
        elif v == "panel":
            K.text(sl, 6.6, 0.45, W - M - 6.6, 0.4, [self.P(self.R(lab, 18, K.MUTED, "strong"))])
            K.text(sl, 6.6, 0.8, W - M - 6.6, 1.05, [self.P(self.R(c["name"], 52, K.ACC, "display"))], anchor="m", name="Title")
            K.text(sl, M, 0.6, 5.6, 1.3, [self.P(self.R(c["definition"], 20, K.MUTED, "body"), align="l", line=1.04)], anchor="m")
            K.rect(sl, 0, 2.1, W, 2.15, fill=rgb(st["panel"]), name="Quote band")
            ptxt = rgb(st.get("panel_text", st["text"]))
            phl = rgb(st.get("panel_hl", st["accent"]))
            K.text(sl, M, 2.1, W - 2 * M, 2.15, self.excerpt_paras(c["quote"], c["hl"], qs, ptxt, phl, align="ctr"), anchor="m")
            K.text(sl, 6.95, 4.6, W - M - 6.95, 2.3, how)
            K.text(sl, M, 4.6, 5.75, 2.3, why)
        self.pageno(sl)

    # =========================== מילים שמעוררות מלחמה ===========================
    def words(self):
        s, st, K = self.s, self.st, self.K
        v = st["v"]["words"]
        sl = self.new(0.6)
        self.title(sl, "מילים שמעוררות את המלחמה")
        K.text(sl, M, 1.42, W - 2 * M, 0.5, [self.P(self.R(s["words_intro"], 20, K.MUTED, "body"))])
        words = s["words"]
        wf = st.get("wfactor", 0.5)
        if v == "scatter":
            # פריסה זורמת: מילים בגדלים שונים, מימין לשמאל, בלי שבירת צירופים
            base_sizes = st.get("scatter_sizes", [62, 38, 44, 36, 56, 34, 40, 50])
            hi = [True, False, False, False, True, False, False, True]
            avail = W - 2 * M
            gap = 0.6
            for scale in (1.0, 0.92, 0.85, 0.78, 0.72, 0.66):
                sizes = [z * scale for z in base_sizes]
                rows = self._flow(words, sizes, wf, avail, gap)
                heights = [max(z for _, _, z, _ in r) / 72 * 1.5 for r in rows]
                total = sum(heights) + 0.15 * (len(rows) - 1)
                if total <= 4.7:
                    break
            y = 2.05 + max(0, (4.75 - total) / 2)
            for r, row in enumerate(rows):
                xr = W - M - [0, 0.5, 0.9][r % 3]
                for k, (i, wd, sz, ww) in enumerate(row):
                    off = [0, 0.16, 0.05][k % 3]
                    K.text(sl, xr - ww, y + off + (heights[r] - sz / 72 * 1.5) / 2, ww, sz / 72 * 1.5,
                           [self.P(self.R(wd.replace(" ", "\u00a0"), sz, K.ACC if hi[i % 8] else K.TEXT,
                                          "display" if hi[i % 8] else "wordfont"))], wrap=False)
                    xr -= ww + gap
                y += heights[r] + 0.15
        elif v == "_old_scatter":
            rows, cur, curw = [], [], 0.0
            for i, wd in enumerate(words):
                sz = sizes[i % len(sizes)]
                ww = est_w(wd, sz, wf) + 0.2
                if ww > avail:
                    sz = sz * avail / ww
                    ww = avail
                if cur and curw + gap + ww > avail - [0, 0.5, 0.9][len(rows) % 3]:
                    rows.append(cur); cur, curw = [], 0.0
                cur.append((i, wd, sz, ww)); curw += (gap if len(cur) > 1 else 0) + ww
            if cur:
                rows.append(cur)
            heights = [max(z for _, _, z, _ in r) / 72 * 1.5 for r in rows]
            total = sum(heights) + 0.15 * (len(rows) - 1)
            y = 2.05 + max(0, (4.75 - total) / 2)
            for r, row in enumerate(rows):
                xr = W - M - [0, 0.5, 0.9][r % 3]
                for k, (i, wd, sz, ww) in enumerate(row):
                    off = [0, 0.16, 0.05][k % 3]
                    K.text(sl, xr - ww, y + off + (heights[r] - sz / 72 * 1.5) / 2, ww, sz / 72 * 1.5,
                           [self.P(self.R(wd.replace(" ", "\u00a0"), sz, K.ACC if hi[i % 8] else K.TEXT,
                                          "display" if hi[i % 8] else "wordfont"))], wrap=False)
                    xr -= ww + gap
                y += heights[r] + 0.15
        elif v == "list":
            colw = (W - 2 * M - 0.8) / 2
            half = (len(words) + 1) // 2
            for i, wd in enumerate(words):
                col, row = (0, i) if i < half else (1, i - half)
                x = W - M - colw - col * (colw + 0.8)
                y = 2.15 + row * 1.12
                if row:
                    K.hair(sl, x, y - 0.12, colw)
                K.text(sl, x + colw - 0.85, y, 0.85, 0.9, [self.P(self.R(f"{i + 1:02d}", 26, K.ACC, "light"), ltr=False)], anchor="m")
                K.text(sl, x, y, colw - 1.0, 0.9, [self.P(self.R(wd, 32, K.TEXT, "wordfont"))], anchor="m")
        elif v == "inline":
            runs = []
            for i, wd in enumerate(words):
                runs.append(self.R(wd.replace(" ", "\u00a0"), 40, K.ACC if i % 2 == 0 else K.TEXT, "display" if i % 2 == 0 else "wordfont"))
                if i < len(words) - 1:
                    runs.append(self.R("   ·   ", 40, K.MUTED, "body"))
            K.text(sl, M + 0.3, 2.1, W - 2 * M - 0.6, 4.6, [self.P(*runs, align="ctr", line=1.45)], anchor="m")
        self.pageno(sl)

    @staticmethod
    def _flow(words, sizes, wf, avail, gap):
        rows, cur, curw = [], [], 0.0
        for i, wd in enumerate(words):
            sz = sizes[i % len(sizes)]
            ww = est_w(wd, sz, wf) + 0.2
            if ww > avail:
                sz, ww = sz * avail / ww, avail
            if cur and curw + gap + ww > avail - [0, 0.5, 0.9][len(rows) % 3]:
                rows.append(cur); cur, curw = [], 0.0
            cur.append((i, wd, sz, ww)); curw += (gap if len(cur) > 1 else 0) + ww
        if cur:
            rows.append(cur)
        return rows

    # =========================== אירועים ===========================
    def events(self):
        s, st, K = self.s, self.st, self.K
        v = st["v"]["events"]
        sl = self.new(0.75)
        photo = prep_photo(st, self.asset(s["event_photo"]))
        focus = s.get("event_focus", 0.4)
        cap = [self.P(self.R(s["event_caption"], 16, K.MUTED, "body"), line=1.02)]
        ev = s["events"]
        if v == "photo-left-rows":
            self.title(sl, "האירועים שהמילים מעוררות")
            picture_fill(sl, photo, M, 1.7, 5.0, 3.5, focus, name="Event photo")
            self.credit(sl, M, 5.24, 5.0, s["event_credit"])
            K.text(sl, M, 5.55, 5.0, 1.3, cap)
            x0, x1 = 6.25, W - M
            y = 1.65
            for i, (lab, t) in enumerate(ev):
                if i:
                    K.hair(sl, x0, y - 0.1, x1 - x0)
                K.text(sl, x1 - 1.6, y, 1.6, 0.9, [self.P(self.R(lab, 18, K.ACC, "strong"), line=1.0)])
                K.text(sl, x0, y, x1 - 1.6 - x0 - 0.2, 0.95, [self.P(self.R(t, 21), line=1.0)])
                y += 1.04
        elif v == "photo-top":
            self.title(sl, "האירועים שהמילים מעוררות")
            picture_fill(sl, photo, 0, 1.45, W, 2.2, focus, name="Event photo")
            K.text(sl, M + 3.7, 3.7, W - 2 * M - 3.7, 0.55, cap, anchor="t")
            self.credit(sl, M, 3.7, 3.5, s["event_credit"])
            colw = (W - 2 * M - 0.6) / 2
            for i, (lab, t) in enumerate(ev):
                col = 0 if i < 3 else 1
                row = i if i < 3 else i - 3
                x = W - M - colw - col * (colw + 0.6)
                y = 4.35 + row * 0.86
                K.text(sl, x, y, colw, 0.84, [self.P(self.R(lab + "  ", 21, K.ACC, "strong"), self.R(t, 21), line=1.0)])
        elif v == "photo-right-tall":
            pw = 4.6
            picture_fill(sl, photo, W - pw, 0, pw, H, focus, name="Event photo")
            self.credit(sl, W - pw + 0.2, 7.12, pw - 0.4, s["event_credit"], RGBColor(0xEE, 0xEE, 0xEE), align="r")
            x1 = W - pw - 0.6
            self.title(sl, "האירועים שהמילים מעוררות", x=M, w=x1 - M, size=st.get("title_size", 42) - 4)
            y = 1.6
            for i, (lab, t) in enumerate(ev):
                K.text(sl, x1 - 0.7, y - 0.05, 0.7, 0.7, [self.P(self.R(str(i + 1), 34, K.ACC, "display"))])
                K.text(sl, M, y, x1 - 0.85 - M, 0.92, [self.P(self.R(lab + "  ", 21, K.ACC, "strong"), self.R(t, 21), line=1.0)])
                y += 0.95
            K.text(sl, M, 6.35, x1 - M, 0.6, cap)
        self.pageno(sl)

    # =========================== סיכום ===========================
    def summary(self):
        s, st, K = self.s, self.st, self.K
        v = st["v"]["summary"]
        lines = s["summary"]
        sizes = st.get("summary_sizes", [36, 40, 34])
        fonts = ["strong", "display", "light"]
        cols = [K.TEXT, K.ACC, K.TEXT]
        if v == "photo-bg":
            sl = self.new(1.0, photo=self.asset(s["summary_photo"]))
            K.text(sl, M, 1.4, W - 2 * M, 4.6, [self.P(self.R(l, sizes[i], cols[i], fonts[i]), align="ctr", after=14, line=1.05)
                                                for i, l in enumerate(lines)], anchor="m", name="Title")
        elif v == "center":
            sl = self.new(1.0, stripes=st.get("stripes_on_cover", False))
            K.text(sl, M, 1.4, W - 2 * M, 4.6, [self.P(self.R(l, sizes[i], cols[i], fonts[i]), align="ctr", after=14, line=1.05)
                                                for i, l in enumerate(lines)], anchor="m", name="Title")
        elif v == "statement":
            sl = self.new(1.0)
            K.text(sl, M, 0.7, W - 2 * M, 0.45, [self.P(self.R("סיכום", 18, K.ACC, "strong"))])
            K.text(sl, M, 1.2, W - 2 * M, 1.4, [self.P(self.R("״", 110, K.ACC, "display"), line=0.8)])
            K.text(sl, M, 2.35, W - 2 * M - 0.6, 4.2, [self.P(self.R(l, sizes[i], cols[i], fonts[i]), after=12, line=1.05)
                                                      for i, l in enumerate(lines)], anchor="t", name="Title")
        self.pageno(sl)

    def build(self, out):
        self.cover()
        self.part_a()
        self.artist()
        self.lyrics()
        self.concept(0)
        self.concept(1)
        self.words()
        self.events()
        self.summary()
        self.total = self.n
        self.prs.save(out)
        return out
